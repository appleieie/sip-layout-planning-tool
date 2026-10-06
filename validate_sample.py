from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SAMPLE = ROOT / "GBB898die.xlsx"
OUTPUT = ROOT / "sample_die_asset.json"


def normalize_die_name(file_name: str) -> str:
    stem = Path(file_name).stem
    return stem[:-3] if stem.lower().endswith("die") else stem


def round3(value: float) -> float:
    return round(float(value), 3)


def build_sample_asset() -> dict:
    workbook = load_workbook(SAMPLE, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    pads = []
    for row in worksheet.iter_rows(values_only=True):
        try:
            label = row[0]
            name = row[1] or ""
            x = float(row[2])
            y = float(row[3])
            width = float(row[4])
            height = float(row[5])
        except Exception:
            continue
        if label in (None, ""):
            continue
        pads.append({
            "label": str(label),
            "name": str(name),
            "x": round3(x),
            "y": round3(y),
            "width": round3(width),
            "height": round3(height),
        })
    if not pads:
        raise RuntimeError("No pad rows found in sample workbook.")
    die_size_x = max(pad["x"] + pad["width"] / 2 for pad in pads)
    die_size_y = max(pad["y"] + pad["height"] / 2 for pad in pads)
    return {
        "format_version": "3.0",
        "unit": "um",
        "die": {
            "diename": normalize_die_name(SAMPLE.name),
            "diesize": {"x": round3(die_size_x), "y": round3(die_size_y)},
        },
        "pads": pads,
    }


if __name__ == "__main__":
    asset = build_sample_asset()
    OUTPUT.write_text(json.dumps(asset, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"sample_die_asset.json written: {OUTPUT}")
    print(f"die={asset['die']['diename']} pads={len(asset['pads'])} size=({asset['die']['diesize']['x']}, {asset['die']['diesize']['y']})")
